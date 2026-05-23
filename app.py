import base64
import html
import io
import json
import re
import sqlite3
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.parse import unquote

import folium
import pandas as pd
import requests
import streamlit as st
from folium.plugins import Fullscreen, MarkerCluster, MeasureControl
from streamlit_folium import st_folium


# =========================================================
# CONFIGURACIÓN GENERAL
# =========================================================

APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
IMG_DIR = DATA_DIR / "evidencias"
UPLOADS_DIR = DATA_DIR / "uploads"
DB_FILE = DATA_DIR / "lonas_supervision.db"

MAPEABLES_CSV = DATA_DIR / "lonas_mapeables.csv"
PENDIENTES_CSV = DATA_DIR / "lonas_pendientes_sin_coordenada.csv"

DATA_DIR.mkdir(parents=True, exist_ok=True)
IMG_DIR.mkdir(parents=True, exist_ok=True)
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

# Paleta Morena con alto contraste.
MORENA_GUINDA = "#7A0026"
MORENA_GUINDA_DARK = "#4A0018"
MORENA_GUINDA_SOFT = "#A51C48"
MORENA_BEIGE = "#F1E3D3"
MORENA_BEIGE_2 = "#FFF8EF"
MORENA_DORADO = "#C69A2D"
MORENA_GREEN = "#007A3D"
TEXT_DARK = "#1F171A"

STATUS_OPTIONS = [
    "Pendiente",
    "Verificado",
    "Corregir ubicación",
    "Retirar/Reponer lona",
    "No localizada",
]

STATUS_COLORS = {
    "Pendiente": MORENA_GUINDA,
    "Verificado": MORENA_GREEN,
    "Corregir ubicación": "#F59E0B",
    "Retirar/Reponer lona": "#B91C1C",
    "No localizada": "#111827",
}

TILE_OPTIONS = {
    "Calles claro": {
        "tiles": "CartoDB positron",
        "attr": "CartoDB / OpenStreetMap",
    },
    "Calles OSM": {
        "tiles": "OpenStreetMap",
        "attr": "OpenStreetMap",
    },
    "Satélite": {
        "tiles": "https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}",
        "attr": "Esri World Imagery",
    },
    "Terreno / relieve": {
        "tiles": "https://{s}.tile.opentopomap.org/{z}/{x}/{y}.png",
        "attr": "Map data: OpenStreetMap contributors, SRTM | Map style: OpenTopoMap",
    },
    "Oscuro": {
        "tiles": "CartoDB dark_matter",
        "attr": "CartoDB / OpenStreetMap",
    },
}

st.set_page_config(
    page_title="Supervisión de Lonas",
    page_icon="📍",
    layout="wide",
    initial_sidebar_state="expanded",
)


# =========================================================
# CONTROL DE ACCESO SENCILLO POR SECRETS
# =========================================================

def get_auth_users() -> Dict[str, dict]:
    """
    Lee usuarios desde Streamlit Secrets.

    Ejemplo en .streamlit/secrets.toml o Streamlit Cloud Secrets:

    [auth.users.admin]
    name = "Administrador"
    role = "admin"
    password = "1234"
    """
    try:
        return dict(st.secrets["auth"]["users"])
    except Exception:
        return {}


def check_login(username: str, password: str) -> Tuple[bool, dict]:
    users = get_auth_users()

    username = str(username or "").strip().lower()
    password = str(password or "").strip()

    if not username or username not in users:
        return False, {}

    user_data = dict(users[username])
    saved_password = str(user_data.get("password", "")).strip()

    if password == saved_password:
        return True, user_data

    return False, {}


def login_required() -> None:
    if st.session_state.get("authenticated", False):
        st.sidebar.markdown("---")
        st.sidebar.success(f"Acceso: {st.session_state.get('auth_name', 'Usuario')}")
        st.sidebar.caption(f"Rol: {st.session_state.get('auth_role', 'usuario')}")

        if st.sidebar.button("Cerrar sesión", use_container_width=True):
            for key in ["authenticated", "auth_user", "auth_name", "auth_role"]:
                st.session_state.pop(key, None)
            st.rerun()

        return

    users = get_auth_users()

    st.markdown(
        """
        <style>
        [data-testid="stSidebar"] {
            display: none;
        }

        header[data-testid="stHeader"] {
            visibility: hidden;
            height: 0rem;
        }

        .stApp {
            background: linear-gradient(180deg, #fffaf2 0%, #fff8ef 60%, #ffffff 100%);
        }

        .block-container {
            max-width: 560px;
            padding-top: 5rem;
        }

        .login-card {
            background: #ffffff;
            border: 2px solid rgba(122,0,38,.22);
            border-radius: 24px;
            padding: 30px;
            box-shadow: 0 18px 42px rgba(74,0,24,.18);
            margin-bottom: 18px;
        }

        .login-title {
            color: #4A0018;
            font-size: 2rem;
            font-weight: 900;
            margin-bottom: 6px;
        }

        .login-subtitle {
            color: #5f5055;
            font-size: .98rem;
            margin-bottom: 4px;
            font-weight: 600;
        }

        div[data-testid="stFormSubmitButton"] button {
            background: #7A0026 !important;
            color: #ffffff !important;
            border-radius: 12px !important;
            font-weight: 900 !important;
            border: 2px solid #4A0018 !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        """
        <div class="login-card">
            <div class="login-title">Acceso restringido</div>
            <div class="login-subtitle">Supervisión de Lonas</div>
            <div>Ingresa con usuario y contraseña autorizados.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    if not users:
        st.error("No hay usuarios configurados en Secrets.")
        st.info(
            "Agrega usuarios en Streamlit Cloud → App → Settings → Secrets, "
            "o en local dentro de .streamlit/secrets.toml."
        )
        st.stop()

    with st.form("login_form"):
        username = st.text_input("Usuario")
        password = st.text_input("Contraseña", type="password")
        submitted = st.form_submit_button("Entrar", use_container_width=True)

        if submitted:
            ok, user_data = check_login(username, password)

            if ok:
                clean_user = str(username).strip().lower()
                st.session_state["authenticated"] = True
                st.session_state["auth_user"] = clean_user
                st.session_state["auth_name"] = user_data.get("name", clean_user)
                st.session_state["auth_role"] = user_data.get("role", "usuario")
                st.rerun()
            else:
                st.error("Usuario o contraseña incorrectos.")

    st.stop()


login_required()


# =========================================================
# ESTILOS GENERALES
# =========================================================

st.markdown(
    f"""
    <style>
    :root {{
        --guinda: {MORENA_GUINDA};
        --guinda-dark: {MORENA_GUINDA_DARK};
        --guinda-soft: {MORENA_GUINDA_SOFT};
        --beige: {MORENA_BEIGE};
        --beige2: {MORENA_BEIGE_2};
        --dorado: {MORENA_DORADO};
        --texto: {TEXT_DARK};
    }}

    header[data-testid="stHeader"] {{
        background: transparent;
        height: 0rem;
        visibility: hidden;
    }}

    #MainMenu, footer {{
        visibility: hidden;
    }}

    .stApp {{
        background: linear-gradient(180deg, #fffaf2 0%, var(--beige2) 45%, #ffffff 100%);
        color: var(--texto);
    }}

    [data-testid="stSidebar"] {{
        background: linear-gradient(180deg, #ffffff 0%, var(--beige) 100%);
        border-right: 2px solid rgba(122, 0, 38, .28);
    }}

    [data-testid="stSidebar"] h1,
    [data-testid="stSidebar"] h2,
    [data-testid="stSidebar"] h3,
    [data-testid="stSidebar"] label {{
        color: var(--guinda-dark) !important;
        font-weight: 800 !important;
    }}

    .block-container {{
        padding-top: 1rem;
        padding-bottom: 4rem;
    }}

    .hero-card {{
        border-radius: 22px;
        padding: 24px 28px;
        margin-bottom: 18px;
        color: white;
        background: linear-gradient(135deg, var(--guinda-dark) 0%, var(--guinda) 55%, var(--guinda-soft) 100%);
        box-shadow: 0 18px 42px rgba(74, 0, 24, .32);
        border: 1px solid rgba(255,255,255,.35);
    }}

    .hero-title {{
        font-size: 2.08rem;
        line-height: 1.08;
        font-weight: 900;
        letter-spacing: -.02em;
        margin: 0 0 8px 0;
        color: #ffffff;
    }}

    .hero-subtitle {{
        font-size: 1rem;
        margin: 0;
        opacity: .96;
        color: #fff8ef;
        font-weight: 500;
    }}

    .kpi-card {{
        background: #ffffff;
        border: 2px solid rgba(122, 0, 38, .18);
        border-radius: 18px;
        padding: 16px 18px;
        box-shadow: 0 10px 25px rgba(74,0,24,.10);
        min-height: 98px;
    }}

    .kpi-label {{
        color: #4A0018;
        font-size: .78rem;
        text-transform: uppercase;
        letter-spacing: .09em;
        font-weight: 900;
        margin-bottom: 8px;
    }}

    .kpi-number {{
        color: var(--guinda);
        font-size: 2.15rem;
        font-weight: 900;
        line-height: 1;
    }}

    .section-card {{
        background: #ffffff;
        border: 2px solid rgba(122, 0, 38, .16);
        border-radius: 18px;
        padding: 16px 18px;
        box-shadow: 0 10px 24px rgba(74,0,24,.09);
        margin-bottom: 12px;
    }}

    .legend-dot {{
        display: inline-block;
        width: 13px;
        height: 13px;
        border-radius: 50%;
        margin-right: 8px;
        vertical-align: middle;
        border: 1px solid rgba(0,0,0,.20);
    }}

    .note-box {{
        background: #fff3df;
        border: 2px solid rgba(198,154,45,.45);
        border-left: 7px solid var(--dorado);
        border-radius: 14px;
        padding: 13px 15px;
        color: #3b2418;
        margin: 10px 0;
        font-weight: 600;
    }}

    .stButton > button,
    .stDownloadButton > button,
    div[data-testid="stFormSubmitButton"] button {{
        border-radius: 12px !important;
        border: 2px solid rgba(122, 0, 38, .30) !important;
        font-weight: 800 !important;
        color: var(--guinda-dark) !important;
        background: #ffffff !important;
    }}

    .stButton > button:hover,
    .stDownloadButton > button:hover {{
        border-color: var(--guinda) !important;
        color: var(--guinda) !important;
    }}

    div[data-testid="stFormSubmitButton"] button,
    .stDownloadButton > button[kind="primary"] {{
        background: var(--guinda) !important;
        color: white !important;
        border-color: var(--guinda-dark) !important;
    }}

    div[data-baseweb="tab-list"] {{
        gap: 8px;
        border-bottom: 2px solid rgba(122,0,38,.28);
        overflow-x: auto;
        white-space: nowrap;
        padding-bottom: 0px;
    }}

    button[data-baseweb="tab"] {{
        background: #ead8cd !important;
        color: #4A0018 !important;
        border-radius: 999px 999px 0 0 !important;
        padding: 10px 18px !important;
        border: 1px solid rgba(122,0,38,.18) !important;
        font-weight: 800 !important;
        opacity: 1 !important;
    }}

    button[data-baseweb="tab"] p {{
        color: #4A0018 !important;
        font-weight: 800 !important;
    }}

    button[data-baseweb="tab"][aria-selected="true"] {{
        background: var(--guinda) !important;
        color: #ffffff !important;
        border-bottom: 3px solid var(--dorado) !important;
        font-weight: 900 !important;
        box-shadow: 0 8px 18px rgba(74,0,24,.22);
    }}

    button[data-baseweb="tab"][aria-selected="true"] p {{
        color: #ffffff !important;
        font-weight: 900 !important;
    }}

    h1, h2, h3, h4 {{
        color: var(--guinda-dark);
        font-weight: 900;
    }}

    iframe {{
        border-radius: 16px;
        border: 2px solid rgba(122,0,38,.24) !important;
        box-shadow: 0 12px 26px rgba(74,0,24,.12);
    }}

    @media screen and (max-width: 768px) {{
        .block-container {{
            padding-left: .85rem;
            padding-right: .85rem;
            padding-top: .6rem;
            padding-bottom: 6rem;
        }}

        .hero-card {{
            padding: 18px 18px;
            border-radius: 18px;
        }}

        .hero-title {{
            font-size: 1.55rem;
        }}

        .hero-subtitle {{
            font-size: .9rem;
        }}

        .kpi-card {{
            min-height: 86px;
            padding: 14px 15px;
            border: 2px solid rgba(122,0,38,.24);
        }}

        .kpi-number {{
            font-size: 2rem;
        }}

        .kpi-label {{
            font-size: .75rem;
            color: var(--guinda-dark);
        }}

        div[data-baseweb="tab-list"] {{
            background: #fff8ef;
            border-radius: 12px 12px 0 0;
            padding-top: 6px;
            padding-left: 4px;
            padding-right: 4px;
        }}

        button[data-baseweb="tab"] {{
            background: #e2c8bc !important;
            color: var(--guinda-dark) !important;
            padding: 10px 16px !important;
            min-width: max-content !important;
        }}

        button[data-baseweb="tab"][aria-selected="true"] {{
            background: var(--guinda) !important;
            color: #ffffff !important;
        }}

        iframe {{
            min-height: 560px !important;
        }}
    }}
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# UTILIDADES GENERALES
# =========================================================

def clean_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value or "").strip()


def to_float(value: object) -> Optional[float]:
    text = clean_text(value)

    if not text:
        return None

    text = text.replace(",", ".")

    try:
        return float(text)
    except Exception:
        return None


def normalize_header(text: object) -> str:
    text = clean_text(text).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"\s+", " ", text)
    return text.replace("\n", " ").strip()


def make_unique_headers(headers: List[object]) -> List[str]:
    counts = {}
    out = []

    for h in headers:
        base = clean_text(h)

        if not base:
            base = "columna"

        if base not in counts:
            counts[base] = 1
            out.append(base)
        else:
            counts[base] += 1
            out.append(f"{base}__{counts[base]}")

    return out


def dataframe_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")


# =========================================================
# BASE DE DATOS SQLITE
# =========================================================

def get_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with get_conn() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS lonas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                unique_key TEXT UNIQUE,
                archivo_origen TEXT,
                hoja_origen TEXT,
                fila_excel INTEGER,
                fecha TEXT,
                responsable TEXT,
                municipio TEXT,
                distrito_local TEXT,
                seccion TEXT,
                colonia TEXT,
                direccion TEXT,
                ciudad_comunidad TEXT,
                nombre_enlace TEXT,
                celular TEXT,
                link_maps TEXT,
                url_expandida TEXT,
                lonas_colocadas REAL,
                fotografia TEXT,
                observaciones TEXT,
                latitud REAL,
                longitud REAL,
                fuente_coordenada TEXT,
                estado_coordenada TEXT,
                estatus TEXT DEFAULT 'Pendiente',
                supervisor TEXT,
                nota_supervision TEXT,
                fecha_revision TEXT,
                fecha_carga TEXT,
                usuario_carga TEXT,
                evidencia_rutas TEXT
            )
            """
        )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS uploads (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                archivo TEXT,
                registros_leidos INTEGER,
                registros_insertados INTEGER,
                registros_duplicados INTEGER,
                registros_mapeables INTEGER,
                registros_pendientes INTEGER,
                fecha_carga TEXT,
                usuario_carga TEXT
            )
            """
        )

        conn.commit()


def db_count() -> int:
    with get_conn() as conn:
        row = conn.execute("SELECT COUNT(*) AS total FROM lonas").fetchone()
        return int(row["total"] or 0)


def create_unique_key(record: dict) -> str:
    parts = [
        clean_text(record.get("archivo_origen")),
        clean_text(record.get("hoja_origen")),
        clean_text(record.get("fila_excel")),
        clean_text(record.get("link_maps")),
        clean_text(record.get("responsable")),
        clean_text(record.get("distrito_local")),
        clean_text(record.get("seccion")),
        clean_text(record.get("direccion")),
    ]

    return "|".join(parts).lower()


def insert_lona(record: dict) -> Tuple[Optional[int], bool]:
    record = dict(record)
    record["unique_key"] = create_unique_key(record)
    record.setdefault("estatus", "Pendiente")
    record.setdefault("fecha_carga", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    record.setdefault("usuario_carga", st.session_state.get("auth_user", ""))

    cols = [
        "unique_key",
        "archivo_origen",
        "hoja_origen",
        "fila_excel",
        "fecha",
        "responsable",
        "municipio",
        "distrito_local",
        "seccion",
        "colonia",
        "direccion",
        "ciudad_comunidad",
        "nombre_enlace",
        "celular",
        "link_maps",
        "url_expandida",
        "lonas_colocadas",
        "fotografia",
        "observaciones",
        "latitud",
        "longitud",
        "fuente_coordenada",
        "estado_coordenada",
        "estatus",
        "supervisor",
        "nota_supervision",
        "fecha_revision",
        "fecha_carga",
        "usuario_carga",
        "evidencia_rutas",
    ]

    values = [record.get(c) for c in cols]

    with get_conn() as conn:
        placeholders = ",".join(["?"] * len(cols))
        sql = f"INSERT OR IGNORE INTO lonas ({','.join(cols)}) VALUES ({placeholders})"
        cur = conn.execute(sql, values)
        conn.commit()

        inserted = cur.rowcount == 1

        row = conn.execute(
            "SELECT id FROM lonas WHERE unique_key = ?",
            (record["unique_key"],),
        ).fetchone()

        return (int(row["id"]) if row else None), inserted


def update_lona_review(
    lona_id: int,
    estatus: str,
    supervisor: str,
    nota: str,
    latitud: Optional[float],
    longitud: Optional[float],
    fuente: str = "captura_manual",
) -> None:
    estado = "exacta" if latitud is not None and longitud is not None else "pendiente"

    with get_conn() as conn:
        conn.execute(
            """
            UPDATE lonas
            SET estatus = ?,
                supervisor = ?,
                nota_supervision = ?,
                latitud = COALESCE(?, latitud),
                longitud = COALESCE(?, longitud),
                fuente_coordenada = CASE WHEN ? IS NOT NULL AND ? IS NOT NULL THEN ? ELSE fuente_coordenada END,
                estado_coordenada = CASE WHEN ? IS NOT NULL AND ? IS NOT NULL THEN ? ELSE estado_coordenada END,
                fecha_revision = ?
            WHERE id = ?
            """,
            (
                estatus,
                supervisor,
                nota,
                latitud,
                longitud,
                latitud,
                longitud,
                fuente,
                latitud,
                longitud,
                estado,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                lona_id,
            ),
        )
        conn.commit()


def update_lona_coords(
    lona_id: int,
    latitud: float,
    longitud: float,
    fuente: str,
    estado: str,
    url_expandida: str = "",
) -> None:
    with get_conn() as conn:
        conn.execute(
            """
            UPDATE lonas
            SET latitud = ?,
                longitud = ?,
                fuente_coordenada = ?,
                estado_coordenada = ?,
                url_expandida = COALESCE(NULLIF(?, ''), url_expandida),
                fecha_revision = ?
            WHERE id = ?
            """,
            (
                latitud,
                longitud,
                fuente,
                estado,
                url_expandida,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                lona_id,
            ),
        )
        conn.commit()


def update_lona_expanded_url(lona_id: int, expanded_url: str) -> None:
    with get_conn() as conn:
        conn.execute(
            "UPDATE lonas SET url_expandida = ? WHERE id = ?",
            (expanded_url, lona_id),
        )
        conn.commit()


def load_lonas_df() -> pd.DataFrame:
    with get_conn() as conn:
        df = pd.read_sql_query("SELECT * FROM lonas ORDER BY id ASC", conn)

    if df.empty:
        return df

    for col in ["latitud", "longitud", "lonas_colocadas"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df["latitud_mapa"] = df["latitud"]
    df["longitud_mapa"] = df["longitud"]
    df["estatus"] = df["estatus"].fillna("Pendiente").replace("", "Pendiente")

    return df


def seed_from_csv_if_empty() -> None:
    if db_count() > 0:
        return

    if MAPEABLES_CSV.exists():
        df_seed = pd.read_csv(MAPEABLES_CSV, dtype=str, encoding="utf-8-sig").fillna("")

        for _, row in df_seed.iterrows():
            lat = to_float(row.get("latitud"))
            lon = to_float(row.get("longitud"))

            record = {
                "archivo_origen": "base_inicial_csv",
                "hoja_origen": "mapeables",
                "fila_excel": int(to_float(row.get("fila_excel")) or 0),
                "fecha": clean_text(row.get("fecha")),
                "responsable": clean_text(row.get("responsable")),
                "municipio": clean_text(row.get("municipio")),
                "distrito_local": clean_text(row.get("distrito_local")),
                "seccion": clean_text(row.get("seccion")),
                "colonia": clean_text(row.get("colonia")),
                "direccion": clean_text(row.get("direccion")),
                "ciudad_comunidad": clean_text(row.get("ciudad_comunidad")),
                "nombre_enlace": clean_text(row.get("nombre_enlace")),
                "celular": clean_text(row.get("celular")),
                "link_maps": clean_text(row.get("link_maps")),
                "url_expandida": clean_text(row.get("link_maps")),
                "lonas_colocadas": to_float(row.get("lonas_colocadas")),
                "fotografia": clean_text(row.get("fotografia")),
                "observaciones": clean_text(row.get("observaciones")),
                "latitud": lat,
                "longitud": lon,
                "fuente_coordenada": "base_inicial",
                "estado_coordenada": "exacta" if lat is not None and lon is not None else "pendiente",
                "estatus": "Pendiente",
                "supervisor": "",
                "nota_supervision": "",
                "fecha_revision": "",
                "evidencia_rutas": "",
            }

            insert_lona(record)

    if PENDIENTES_CSV.exists():
        df_p = pd.read_csv(PENDIENTES_CSV, dtype=str, encoding="utf-8-sig").fillna("")

        for _, row in df_p.iterrows():
            record = {
                "archivo_origen": "base_inicial_csv",
                "hoja_origen": "pendientes",
                "fila_excel": int(to_float(row.get("fila_excel")) or 0),
                "fecha": clean_text(row.get("fecha")),
                "responsable": clean_text(row.get("responsable")),
                "municipio": clean_text(row.get("municipio")),
                "distrito_local": clean_text(row.get("distrito_local")),
                "seccion": clean_text(row.get("seccion")),
                "colonia": clean_text(row.get("colonia")),
                "direccion": clean_text(row.get("direccion")),
                "ciudad_comunidad": clean_text(row.get("ciudad_comunidad")),
                "nombre_enlace": clean_text(row.get("nombre_enlace")),
                "celular": clean_text(row.get("celular")),
                "link_maps": clean_text(row.get("link_maps")),
                "url_expandida": "",
                "lonas_colocadas": to_float(row.get("lonas_colocadas")),
                "fotografia": clean_text(row.get("fotografia")),
                "observaciones": clean_text(row.get("observaciones")),
                "latitud": None,
                "longitud": None,
                "fuente_coordenada": "link_no_resuelto",
                "estado_coordenada": "pendiente",
                "estatus": "Pendiente",
                "supervisor": "",
                "nota_supervision": "",
                "fecha_revision": "",
                "evidencia_rutas": "",
            }

            insert_lona(record)


# =========================================================
# GOOGLE MAPS / COORDENADAS
# =========================================================

def extract_coords_from_google_url(url: str) -> Tuple[Optional[float], Optional[float]]:
    if not url:
        return None, None

    decoded = unquote(str(url))

    patterns = [
        r"@(-?\d+\.\d+),(-?\d+\.\d+)",
        r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)",
        r"[?&]q=(-?\d+\.\d+),(-?\d+\.\d+)",
        r"[?&]ll=(-?\d+\.\d+),(-?\d+\.\d+)",
        r"query=(-?\d+\.\d+),(-?\d+\.\d+)",
        r"destination=(-?\d+\.\d+),(-?\d+\.\d+)",
    ]

    for pattern in patterns:
        match = re.search(pattern, decoded)

        if match:
            lat = float(match.group(1))
            lon = float(match.group(2))

            if -90 <= lat <= 90 and -180 <= lon <= 180:
                return lat, lon

    return None, None


def expand_google_maps_url(url: str) -> str:
    url = clean_text(url)

    if not url:
        return ""

    try:
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0 Safari/537.36"
            )
        }

        response = requests.get(
            url,
            headers=headers,
            allow_redirects=True,
            timeout=12,
        )

        return response.url or url

    except Exception:
        return url


def resolve_maps_link(url: str, try_expand: bool = True) -> dict:
    original_url = clean_text(url)

    if not original_url:
        return {
            "latitud": None,
            "longitud": None,
            "url_expandida": "",
            "fuente_coordenada": "sin_link",
            "estado_coordenada": "pendiente",
        }

    lat, lon = extract_coords_from_google_url(original_url)

    if lat is not None and lon is not None:
        return {
            "latitud": lat,
            "longitud": lon,
            "url_expandida": original_url,
            "fuente_coordenada": "link_directo",
            "estado_coordenada": "exacta",
        }

    if not try_expand:
        return {
            "latitud": None,
            "longitud": None,
            "url_expandida": "",
            "fuente_coordenada": "link_no_resuelto",
            "estado_coordenada": "pendiente",
        }

    expanded = expand_google_maps_url(original_url)
    lat, lon = extract_coords_from_google_url(expanded)

    if lat is not None and lon is not None:
        return {
            "latitud": lat,
            "longitud": lon,
            "url_expandida": expanded,
            "fuente_coordenada": "link_corto_resuelto",
            "estado_coordenada": "exacta",
        }

    return {
        "latitud": None,
        "longitud": None,
        "url_expandida": expanded,
        "fuente_coordenada": "link_no_resuelto",
        "estado_coordenada": "pendiente",
    }


# =========================================================
# LECTURA DE EXCEL
# =========================================================

def detect_header_row(raw: pd.DataFrame) -> Optional[int]:
    max_scan = min(len(raw), 15)

    for i in range(max_scan):
        row_values = [normalize_header(x) for x in raw.iloc[i].tolist()]
        joined = " | ".join(row_values)

        keywords = [
            "fecha",
            "responsable",
            "municipio",
            "distrito",
            "seccion",
            "colonia",
            "direccion",
            "link",
            "maps",
            "lona",
            "observacion",
        ]

        score = sum(1 for kw in keywords if kw in joined)

        if score >= 4 and "fecha" in joined:
            return i

    return None


def build_column_groups(columns: List[str]) -> Dict[str, List[str]]:
    groups = {
        "fecha": [],
        "responsable": [],
        "municipio": [],
        "distrito_local": [],
        "seccion": [],
        "colonia": [],
        "direccion": [],
        "ciudad_comunidad": [],
        "nombre_enlace": [],
        "celular": [],
        "link_maps": [],
        "lonas_colocadas": [],
        "fotografia": [],
        "observaciones": [],
    }

    for col in columns:
        norm = normalize_header(col)

        if "fecha" in norm:
            groups["fecha"].append(col)
        elif "responsable" in norm:
            groups["responsable"].append(col)
        elif "municipio" in norm:
            groups["municipio"].append(col)
        elif "distrito" in norm:
            groups["distrito_local"].append(col)
        elif "seccion" in norm:
            groups["seccion"].append(col)
        elif "colonia" in norm:
            groups["colonia"].append(col)
        elif "direccion" in norm:
            groups["direccion"].append(col)
        elif "ciudad" in norm or "comunidad" in norm:
            groups["ciudad_comunidad"].append(col)
        elif "enlace" in norm and "link" not in norm:
            groups["nombre_enlace"].append(col)
        elif "celular" in norm or "telefono" in norm:
            groups["celular"].append(col)
        elif "link" in norm or "maps" in norm or "mapa" in norm or "ubicacion" in norm:
            groups["link_maps"].append(col)
        elif "lona" in norm and ("colocada" in norm or "colocadas" in norm or "cantidad" in norm):
            groups["lonas_colocadas"].append(col)
        elif "foto" in norm or "fotografia" in norm or "evidencia" in norm:
            groups["fotografia"].append(col)
        elif "observacion" in norm or "comentario" in norm:
            groups["observaciones"].append(col)

    return groups


def get_first_value(row: pd.Series, cols: List[str], prefer_last: bool = False) -> str:
    if not cols:
        return ""

    iterable = list(reversed(cols)) if prefer_last else cols

    for col in iterable:
        if col in row.index:
            value = clean_text(row[col])

            if value:
                return value

    return ""


def read_excel_images(excel_bytes: bytes, sheet_name: str) -> Dict[int, List[Tuple[bytes, str]]]:
    images_by_row: Dict[int, List[Tuple[bytes, str]]] = {}

    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
        ws = wb[sheet_name]

        for image in getattr(ws, "_images", []):
            try:
                anchor_row = int(image.anchor._from.row) + 1
                img_bytes = image._data()
                ext = str(getattr(image, "format", "png") or "png").lower()
                images_by_row.setdefault(anchor_row, []).append((img_bytes, ext))
            except Exception:
                continue

    except Exception:
        return {}

    return images_by_row


def save_images_for_lona(lona_id: int, images: List[Tuple[bytes, str]]) -> List[str]:
    saved = []

    for idx, (img_bytes, ext) in enumerate(images, start=1):
        ext = ext.lower().replace(".", "")

        if ext not in ["png", "jpg", "jpeg", "webp"]:
            ext = "png"

        path = IMG_DIR / f"db_{lona_id}_evidencia_{idx}.{ext}"

        try:
            path.write_bytes(img_bytes)
            saved.append(str(path.relative_to(APP_DIR)))
        except Exception:
            continue

    if saved:
        with get_conn() as conn:
            conn.execute(
                "UPDATE lonas SET evidencia_rutas = ? WHERE id = ?",
                (json.dumps(saved, ensure_ascii=False), lona_id),
            )
            conn.commit()

    return saved


def parse_excel_file(uploaded_file, resolve_short_links: bool) -> Tuple[int, int, int, int, int]:
    excel_bytes = uploaded_file.getvalue()
    archivo_origen = uploaded_file.name

    uploaded_copy = UPLOADS_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{archivo_origen}"
    uploaded_copy.write_bytes(excel_bytes)

    total_records = 0
    inserted = 0
    duplicated = 0
    mapeables = 0
    pendientes = 0

    xls = pd.ExcelFile(io.BytesIO(excel_bytes))

    for sheet_name in xls.sheet_names:
        raw = pd.read_excel(
            io.BytesIO(excel_bytes),
            sheet_name=sheet_name,
            header=None,
            dtype=str,
        ).fillna("")

        header_idx = detect_header_row(raw)

        if header_idx is None:
            continue

        headers = make_unique_headers(raw.iloc[header_idx].tolist())
        data = raw.iloc[header_idx + 1:].copy()
        data.columns = headers

        groups = build_column_groups(headers)
        images_by_row = read_excel_images(excel_bytes, sheet_name)

        for raw_idx, row in data.iterrows():
            excel_row_number = int(raw_idx) + 1

            record = {
                "archivo_origen": archivo_origen,
                "hoja_origen": sheet_name,
                "fila_excel": excel_row_number,
                "fecha": get_first_value(row, groups["fecha"]),
                "responsable": get_first_value(row, groups["responsable"]),
                "municipio": get_first_value(row, groups["municipio"]),
                "distrito_local": get_first_value(row, groups["distrito_local"]),
                "seccion": get_first_value(row, groups["seccion"]),
                "colonia": get_first_value(row, groups["colonia"]),
                "direccion": get_first_value(row, groups["direccion"]),
                "ciudad_comunidad": get_first_value(row, groups["ciudad_comunidad"]),
                "nombre_enlace": get_first_value(row, groups["nombre_enlace"]),
                "celular": get_first_value(row, groups["celular"], prefer_last=True),
                "link_maps": get_first_value(row, groups["link_maps"]),
                "lonas_colocadas": to_float(get_first_value(row, groups["lonas_colocadas"])) or 1,
                "fotografia": get_first_value(row, groups["fotografia"]),
                "observaciones": get_first_value(row, groups["observaciones"]),
                "estatus": "Pendiente",
                "supervisor": "",
                "nota_supervision": "",
                "fecha_revision": "",
                "fecha_carga": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "usuario_carga": st.session_state.get("auth_user", ""),
                "evidencia_rutas": "",
            }

            meaningful = [
                record["fecha"],
                record["responsable"],
                record["municipio"],
                record["distrito_local"],
                record["seccion"],
                record["colonia"],
                record["direccion"],
                record["link_maps"],
                record["observaciones"],
            ]

            if not any(clean_text(x) for x in meaningful):
                continue

            coord = resolve_maps_link(record["link_maps"], try_expand=resolve_short_links)
            record.update(coord)

            total_records += 1

            if record.get("latitud") is not None and record.get("longitud") is not None:
                mapeables += 1
            else:
                pendientes += 1

            lona_id, was_inserted = insert_lona(record)

            if was_inserted:
                inserted += 1

                if lona_id and excel_row_number in images_by_row:
                    save_images_for_lona(lona_id, images_by_row[excel_row_number])
            else:
                duplicated += 1

    with get_conn() as conn:
        conn.execute(
            """
            INSERT INTO uploads (
                archivo,
                registros_leidos,
                registros_insertados,
                registros_duplicados,
                registros_mapeables,
                registros_pendientes,
                fecha_carga,
                usuario_carga
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                archivo_origen,
                total_records,
                inserted,
                duplicated,
                mapeables,
                pendientes,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                st.session_state.get("auth_user", ""),
            ),
        )
        conn.commit()

    return total_records, inserted, duplicated, mapeables, pendientes


# =========================================================
# IMÁGENES / MAPAS / EXPORTACIONES
# =========================================================

def image_files_for_lona(row: pd.Series) -> List[Path]:
    files: List[Path] = []

    evidencia_rutas = clean_text(row.get("evidencia_rutas", ""))

    if evidencia_rutas:
        try:
            rutas = json.loads(evidencia_rutas)

            for ruta in rutas:
                path = APP_DIR / ruta

                if path.exists():
                    files.append(path)
        except Exception:
            pass

    lona_id = row.get("id")

    if pd.notna(lona_id):
        files.extend(sorted(IMG_DIR.glob(f"db_{int(lona_id)}_evidencia_*.*")))

    fila_excel = row.get("fila_excel")

    if pd.notna(fila_excel):
        files.extend(sorted(IMG_DIR.glob(f"fila_{int(fila_excel)}_evidencia_*.*")))

    unique = []
    seen = set()

    for f in files:
        if f.exists() and str(f) not in seen:
            unique.append(f)
            seen.add(str(f))

    return unique


def img_to_base64(path: Path) -> Optional[str]:
    try:
        data = path.read_bytes()
        ext = path.suffix.lower().lstrip(".") or "jpg"

        if ext == "jpg":
            ext = "jpeg"

        return f"data:image/{ext};base64," + base64.b64encode(data).decode("ascii")

    except Exception:
        return None


def make_popup_html(row: pd.Series, include_img: bool = True) -> str:
    lona_id = int(row.get("id", 0))
    img_html = ""

    if include_img:
        imgs = image_files_for_lona(row)

        if imgs:
            src = img_to_base64(imgs[0])

            if src:
                img_html = f"""
                <div style='margin-top:10px'>
                    <img src='{src}' style='max-width:260px; max-height:190px; border-radius:10px; border:1px solid #e5d8cf;'>
                </div>
                """

    maps_link = html.escape(clean_text(row.get("link_maps", "")))

    link_html = (
        f"<a href='{maps_link}' target='_blank' style='color:{MORENA_GUINDA};font-weight:bold'>Abrir en Google Maps</a>"
        if maps_link
        else ""
    )

    status = html.escape(clean_text(row.get("estatus", "Pendiente")) or "Pendiente")
    status_color = STATUS_COLORS.get(status, MORENA_GUINDA)

    body = f"""
    <div style='font-family:Arial; width:300px; color:#272124'>
      <h4 style='margin:0 0 8px 0; color:{MORENA_GUINDA_DARK}'>Lona | ID {lona_id}</h4>
      <div style='display:inline-block;background:{status_color};color:white;padding:4px 9px;border-radius:999px;font-size:12px;font-weight:bold;margin-bottom:8px'>{status}</div><br>
      <b>Distrito:</b> {html.escape(clean_text(row.get('distrito_local','')))} &nbsp; <b>Sección:</b> {html.escape(clean_text(row.get('seccion','')))}<br>
      <b>Colonia:</b> {html.escape(clean_text(row.get('colonia','')))}<br>
      <b>Dirección:</b> {html.escape(clean_text(row.get('direccion','')))}<br>
      <b>Responsable:</b> {html.escape(clean_text(row.get('responsable','')))}<br>
      <b>Observaciones:</b> {html.escape(clean_text(row.get('observaciones','')))}<br>
      <b>Supervisor:</b> {html.escape(clean_text(row.get('supervisor','')))}<br>
      <b>Nota:</b> {html.escape(clean_text(row.get('nota_supervision','')))}<br>
      <b>Fuente coord.:</b> {html.escape(clean_text(row.get('fuente_coordenada','')))}<br>
      {link_html}
      {img_html}
    </div>
    """

    return body


def add_tile_layers(m: folium.Map, selected_tile: str) -> None:
    selected_tile = selected_tile if selected_tile in TILE_OPTIONS else "Calles claro"

    for name, cfg in TILE_OPTIONS.items():
        folium.TileLayer(
            tiles=cfg["tiles"],
            attr=cfg["attr"],
            name=name,
            overlay=False,
            control=True,
            show=(name == selected_tile),
        ).add_to(m)


def make_map(
    df: pd.DataFrame,
    selected_lona_id: Optional[int] = None,
    selected_tile: str = "Calles claro",
    cluster_points: bool = False,
) -> folium.Map:
    valid = df.dropna(subset=["latitud_mapa", "longitud_mapa"]).copy()

    if valid.empty:
        m = folium.Map(location=[25.79, -109.0], zoom_start=12, tiles=None, control_scale=True)
        add_tile_layers(m, selected_tile)
        folium.LayerControl(collapsed=False).add_to(m)
        return m

    center = [valid["latitud_mapa"].mean(), valid["longitud_mapa"].mean()]
    m = folium.Map(location=center, zoom_start=13, tiles=None, control_scale=True)

    add_tile_layers(m, selected_tile)

    if cluster_points:
        target_layer = MarkerCluster(name="Lonas agrupadas").add_to(m)
    else:
        target_layer = folium.FeatureGroup(name="Lonas", show=True).add_to(m)

    for _, row in valid.iterrows():
        lona_id = int(row["id"])
        status = clean_text(row.get("estatus", "Pendiente")) or "Pendiente"
        color = STATUS_COLORS.get(status, MORENA_GUINDA)
        radius = 12 if selected_lona_id == lona_id else 7

        folium.CircleMarker(
            location=[float(row["latitud_mapa"]), float(row["longitud_mapa"])],
            radius=radius,
            popup=folium.Popup(make_popup_html(row), max_width=370),
            tooltip=f"ID {lona_id} | D{row.get('distrito_local','')} S{row.get('seccion','')} | {status}",
            color="#FFFFFF" if selected_lona_id == lona_id else color,
            weight=3 if selected_lona_id == lona_id else 2,
            fill=True,
            fill_color=color,
            fill_opacity=0.90,
        ).add_to(target_layer)

    if selected_lona_id:
        sel = valid[valid["id"].astype(int) == int(selected_lona_id)]

        if not sel.empty:
            row = sel.iloc[0]

            folium.Marker(
                location=[float(row["latitud_mapa"]), float(row["longitud_mapa"])],
                popup=folium.Popup(make_popup_html(row), max_width=370),
                tooltip="Registro seleccionado",
                icon=folium.Icon(color="darkred", icon="star"),
            ).add_to(m)

    Fullscreen().add_to(m)
    MeasureControl(primary_length_unit="meters", secondary_length_unit="kilometers").add_to(m)
    folium.LayerControl(collapsed=True).add_to(m)

    return m


def make_supervision_cluster_map(
    df: pd.DataFrame,
    selected_tile: str = "Satélite",
) -> folium.Map:
    valid = df.dropna(subset=["latitud_mapa", "longitud_mapa"]).copy()

    if valid.empty:
        m = folium.Map(location=[25.79, -109.0], zoom_start=8, tiles=None, control_scale=True)
        add_tile_layers(m, selected_tile)
        Fullscreen().add_to(m)
        folium.LayerControl(collapsed=True).add_to(m)
        return m

    center = [valid["latitud_mapa"].mean(), valid["longitud_mapa"].mean()]
    m = folium.Map(location=center, zoom_start=8, tiles=None, control_scale=True)

    add_tile_layers(m, selected_tile)

    icon_create_function = f"""
    function(cluster) {{
        var count = cluster.getChildCount();
        var size = count < 10 ? 42 : count < 50 ? 50 : 60;
        return new L.DivIcon({{
            html: '<div style="background:{MORENA_GUINDA}; color:white; width:' + size + 'px; height:' + size + 'px; line-height:' + size + 'px; border-radius:50%; text-align:center; font-weight:900; border:4px solid {MORENA_DORADO}; box-shadow:0 8px 20px rgba(74,0,24,.42); font-size:17px;">' + count + '</div>',
            className: 'marker-cluster-morena',
            iconSize: new L.Point(size, size)
        }});
    }}
    """

    cluster = MarkerCluster(
        name="Total de lonas por distribución geográfica",
        icon_create_function=icon_create_function,
        options={
            "spiderfyOnMaxZoom": True,
            "showCoverageOnHover": True,
            "zoomToBoundsOnClick": True,
            "disableClusteringAtZoom": 18,
            "maxClusterRadius": 80,
        },
    ).add_to(m)

    for _, row in valid.iterrows():
        lona_id = int(row["id"])
        status = clean_text(row.get("estatus", "Pendiente")) or "Pendiente"
        color = STATUS_COLORS.get(status, MORENA_GUINDA)

        tooltip = (
            f"ID {lona_id} | Distrito {row.get('distrito_local','')} | "
            f"Sección {row.get('seccion','')} | {status}"
        )

        folium.CircleMarker(
            location=[float(row["latitud_mapa"]), float(row["longitud_mapa"])],
            radius=7,
            popup=folium.Popup(make_popup_html(row), max_width=370),
            tooltip=tooltip,
            color="#FFFFFF",
            weight=2,
            fill=True,
            fill_color=color,
            fill_opacity=0.92,
        ).add_to(cluster)

    Fullscreen().add_to(m)
    MeasureControl(primary_length_unit="meters", secondary_length_unit="kilometers").add_to(m)
    folium.LayerControl(collapsed=True).add_to(m)

    return m


def kml_escape(value: object) -> str:
    return html.escape(str(value or ""), quote=True)


def hex_to_kml_color(hex_color: str, alpha: str = "ff") -> str:
    color = hex_color.lstrip("#")

    if len(color) != 6:
        color = "7A0026"

    rr, gg, bb = color[0:2], color[2:4], color[4:6]

    return f"{alpha}{bb}{gg}{rr}"


def build_kmz_bytes(df: pd.DataFrame, filename_prefix: str = "lonas_supervision") -> bytes:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    valid = df.dropna(subset=["latitud_mapa", "longitud_mapa"]).copy()

    kml_parts = [
        "<?xml version='1.0' encoding='UTF-8'?>",
        "<kml xmlns='http://www.opengis.net/kml/2.2'>",
        "<Document>",
        f"<name>{filename_prefix}</name>",
        f"<description>KMZ generado desde la app de supervisión. Fecha: {kml_escape(now)}</description>",
        f"<Style id='pendiente'><IconStyle><color>{hex_to_kml_color(MORENA_GUINDA)}</color><scale>1.1</scale><Icon><href>http://maps.google.com/mapfiles/kml/pushpin/red-pushpin.png</href></Icon></IconStyle></Style>",
        f"<Style id='verificado'><IconStyle><color>{hex_to_kml_color(MORENA_GREEN)}</color><scale>1.1</scale><Icon><href>http://maps.google.com/mapfiles/kml/pushpin/grn-pushpin.png</href></Icon></IconStyle></Style>",
        f"<Style id='alerta'><IconStyle><color>{hex_to_kml_color('#B91C1C')}</color><scale>1.1</scale><Icon><href>http://maps.google.com/mapfiles/kml/pushpin/ylw-pushpin.png</href></Icon></IconStyle></Style>",
    ]

    for distrito, df_d in valid.groupby(valid["distrito_local"].astype(str), dropna=False):
        kml_parts.append(f"<Folder><name>Distrito Local {kml_escape(distrito)}</name>")

        for seccion, df_s in df_d.groupby(df_d["seccion"].astype(str), dropna=False):
            kml_parts.append(f"<Folder><name>Sección {kml_escape(seccion)}</name>")

            for _, row in df_s.iterrows():
                lona_id = int(row["id"])
                status = clean_text(row.get("estatus", "Pendiente")) or "Pendiente"

                if status == "Verificado":
                    style = "verificado"
                elif status in ["Pendiente", ""]:
                    style = "pendiente"
                else:
                    style = "alerta"

                imgs = image_files_for_lona(row)
                img_html = ""

                if imgs:
                    img_path = f"files/{imgs[0].name}"
                    img_html = f"<br/><br/><img src='{img_path}' width='420'/>"

                desc = f"""
                <![CDATA[
                <div style='font-family:Arial'>
                  <h3>Lona | ID {lona_id}</h3>
                  <b>Estatus:</b> {html.escape(status)}<br/>
                  <b>Distrito Local:</b> {html.escape(clean_text(row.get('distrito_local','')))}<br/>
                  <b>Sección:</b> {html.escape(clean_text(row.get('seccion','')))}<br/>
                  <b>Municipio:</b> {html.escape(clean_text(row.get('municipio','')))}<br/>
                  <b>Colonia:</b> {html.escape(clean_text(row.get('colonia','')))}<br/>
                  <b>Dirección:</b> {html.escape(clean_text(row.get('direccion','')))}<br/>
                  <b>Responsable:</b> {html.escape(clean_text(row.get('responsable','')))}<br/>
                  <b>Observaciones:</b> {html.escape(clean_text(row.get('observaciones','')))}<br/>
                  <b>Supervisor:</b> {html.escape(clean_text(row.get('supervisor','')))}<br/>
                  <b>Nota supervisión:</b> {html.escape(clean_text(row.get('nota_supervision','')))}<br/>
                  <b>Fecha revisión:</b> {html.escape(clean_text(row.get('fecha_revision','')))}<br/>
                  <b>Fuente coordenada:</b> {html.escape(clean_text(row.get('fuente_coordenada','')))}<br/>
                  {img_html}
                </div>
                ]]>
                """

                name = f"ID {lona_id} | D{row.get('distrito_local','')} S{row.get('seccion','')} | {status}"

                kml_parts.extend(
                    [
                        "<Placemark>",
                        f"<name>{kml_escape(name)}</name>",
                        f"<styleUrl>#{style}</styleUrl>",
                        f"<description>{desc}</description>",
                        "<Point>",
                        f"<coordinates>{float(row['longitud_mapa'])},{float(row['latitud_mapa'])},0</coordinates>",
                        "</Point>",
                        "</Placemark>",
                    ]
                )

            kml_parts.append("</Folder>")

        kml_parts.append("</Folder>")

    kml_parts.extend(["</Document>", "</kml>"])
    kml_text = "\n".join(kml_parts)

    buff = io.BytesIO()

    with zipfile.ZipFile(buff, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("doc.kml", kml_text.encode("utf-8"))

        added = set()

        for _, row in valid.iterrows():
            for img in image_files_for_lona(row):
                arc = f"files/{img.name}"

                if arc not in added and img.exists():
                    zf.write(img, arc)
                    added.add(arc)

    return buff.getvalue()


def render_kpi(label: str, value: object, note: str = "") -> None:
    note_html = (
        f"<div style='font-size:.77rem;color:#4A0018;margin-top:6px;font-weight:600'>{html.escape(str(note))}</div>"
        if note
        else ""
    )

    st.markdown(
        f"""
        <div class='kpi-card'>
            <div class='kpi-label'>{html.escape(str(label))}</div>
            <div class='kpi-number'>{html.escape(str(value))}</div>
            {note_html}
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_status_legend() -> None:
    rows = []

    for status in STATUS_OPTIONS:
        rows.append(
            f"<div style='margin:7px 0;font-weight:700;color:#1F171A'><span class='legend-dot' style='background:{STATUS_COLORS[status]}'></span>{html.escape(status)}</div>"
        )

    st.markdown("".join(rows), unsafe_allow_html=True)


def filter_df(
    df: pd.DataFrame,
    distritos,
    secciones,
    estatuses,
    query: str,
) -> pd.DataFrame:
    if df.empty:
        return df.copy()

    out = df.copy()

    if distritos:
        out = out[out["distrito_local"].astype(str).isin([str(x) for x in distritos])]

    if secciones:
        out = out[out["seccion"].astype(str).isin([str(x) for x in secciones])]

    if estatuses:
        out = out[out["estatus"].isin(estatuses)]

    q = clean_text(query).lower()

    if q:
        cols = [
            c
            for c in [
                "archivo_origen",
                "responsable",
                "municipio",
                "distrito_local",
                "seccion",
                "colonia",
                "direccion",
                "nombre_enlace",
                "celular",
                "observaciones",
                "link_maps",
            ]
            if c in out.columns
        ]

        mask = pd.Series(False, index=out.index)

        for col in cols:
            mask = mask | out[col].astype(str).str.lower().str.contains(
                re.escape(q),
                na=False,
                regex=True,
            )

        out = out[mask]

    return out


# =========================================================
# INICIALIZACIÓN
# =========================================================

init_db()
seed_from_csv_if_empty()
df = load_lonas_df()

st.markdown(
    """
    <div class='hero-card'>
        <div class='hero-title'>📍 Supervisión de Lonas</div>
        <p class='hero-subtitle'>Mapa operativo para validar ubicación, evidencia fotográfica y estatus de revisión.</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.markdown("### Filtros")

if df.empty:
    all_distritos = []
else:
    all_distritos = sorted(
        [x for x in df["distrito_local"].dropna().astype(str).unique().tolist() if x.strip()],
        key=lambda x: (len(x), x),
    )

sel_distritos = st.sidebar.multiselect(
    "Distrito local",
    all_distritos,
    default=all_distritos,
)

seccion_source = df[df["distrito_local"].astype(str).isin(sel_distritos)] if sel_distritos and not df.empty else df

if seccion_source.empty:
    all_secciones = []
else:
    all_secciones = sorted(
        [x for x in seccion_source["seccion"].dropna().astype(str).unique().tolist() if x.strip()],
        key=lambda x: (len(x), x),
    )

sel_secciones = st.sidebar.multiselect("Sección", all_secciones)
sel_estatus = st.sidebar.multiselect("Estatus", STATUS_OPTIONS, default=[])
query = st.sidebar.text_input("Buscar", placeholder="Colonia, dirección, sección...")

st.sidebar.divider()
st.sidebar.markdown("### Visualización del mapa")

map_style = st.sidebar.selectbox(
    "Tipo de mapa base",
    list(TILE_OPTIONS.keys()),
    index=0,
)

cluster_points = st.sidebar.checkbox("Agrupar puntos cercanos", value=False)

filtered = filter_df(df, sel_distritos, sel_secciones, sel_estatus, query)

st.sidebar.divider()
st.sidebar.markdown("### Acciones rápidas")

st.sidebar.download_button(
    "Descargar revisión filtrada CSV",
    dataframe_to_csv_bytes(filtered),
    file_name="revision_lonas_filtrada.csv",
    mime="text/csv",
    use_container_width=True,
)


# =========================================================
# KPIS
# =========================================================

total_registros = len(df)
total_mapeables = int(df["latitud_mapa"].notna().sum()) if not df.empty else 0
total_pendientes_coord = int(df["latitud_mapa"].isna().sum()) if not df.empty else 0
total_verificados = int((df["estatus"] == "Verificado").sum()) if not df.empty else 0
total_pendientes_revision = int((df["estatus"] == "Pendiente").sum()) if not df.empty else 0

k1, k2, k3, k4, k5 = st.columns(5)

with k1:
    render_kpi("Registros", total_registros, "base SQLite")

with k2:
    render_kpi("Mapeables", total_mapeables, "con coordenada")

with k3:
    render_kpi("Verificados", total_verificados, "validación positiva")

with k4:
    render_kpi("Por revisar", total_pendientes_revision, "estatus pendiente")

with k5:
    render_kpi("Sin coordenada", total_pendientes_coord, "requieren captura")

st.write("")


# =========================================================
# PESTAÑAS
# =========================================================

tab_map, tab_supervision_map, tab_summary, tab_review, tab_table, tab_pending, tab_export, tab_help, tab_upload = st.tabs(
    [
        "Mapa",
        "Mapa de supervisión",
        "Resumen",
        "Supervisión",
        "Tabla de supervisión",
        "Pendientes sin coordenada",
        "Exportar",
        "Guía rápida",
        "Cargar Excel",
    ]
)


with tab_map:
    c1, c2 = st.columns([3.25, 1])

    with c2:
        st.markdown("<div class='section-card'>", unsafe_allow_html=True)
        st.subheader("Leyenda")
        render_status_legend()
        st.caption(f"Mapa base activo: {map_style}")
        st.markdown("</div>", unsafe_allow_html=True)

        st.markdown("<div class='section-card'>", unsafe_allow_html=True)
        st.write("**Distribución por estatus**")

        mapped_filtered = filtered.dropna(subset=["latitud_mapa", "longitud_mapa"]) if not filtered.empty else filtered

        if mapped_filtered.empty:
            st.caption("Sin registros mapeables con los filtros actuales.")
        else:
            st.dataframe(
                mapped_filtered["estatus"].value_counts().rename_axis("estatus").reset_index(name="total"),
                hide_index=True,
                use_container_width=True,
            )

        st.markdown("</div>", unsafe_allow_html=True)

        st.markdown(
            "<div class='note-box'>Puedes cambiar el mapa base desde el panel izquierdo: calles, satélite, terreno o modo oscuro.</div>",
            unsafe_allow_html=True,
        )

    with c1:
        selected_for_map = st.session_state.get("selected_lona_id")
        m = make_map(
            filtered,
            selected_for_map,
            selected_tile=map_style,
            cluster_points=cluster_points,
        )
        st_folium(m, width=1280, height=680, returned_objects=[])


with tab_supervision_map:
    st.subheader("Mapa de supervisión")
    st.caption(
        "Vista general agrupada: cada burbuja muestra el total de lonas/registros en esa zona. "
        "Al acercarte con zoom, los puntos se despliegan individualmente."
    )

    ms1, ms2 = st.columns([3.25, 1])

    with ms2:
        st.markdown("<div class='section-card'>", unsafe_allow_html=True)
        st.write("**Configuración**")

        supervision_tile = st.selectbox(
            "Mapa base",
            list(TILE_OPTIONS.keys()),
            index=list(TILE_OPTIONS.keys()).index(map_style) if map_style in TILE_OPTIONS else 0,
            key="supervision_tile",
        )

        st.caption("El cluster muestra el total. Click sobre la burbuja para acercar y separar puntos.")
        st.markdown("</div>", unsafe_allow_html=True)

        st.markdown("<div class='section-card'>", unsafe_allow_html=True)
        st.write("**Totales en filtro**")

        filtered_mapped = filtered.dropna(subset=["latitud_mapa", "longitud_mapa"]) if not filtered.empty else filtered
        render_kpi("Registros agrupados", len(filtered_mapped), "según filtros activos")

        if not filtered.empty:
            resumen_dist = (
                filtered.groupby("distrito_local", dropna=False)
                .size()
                .reset_index(name="total")
                .sort_values("total", ascending=False)
            )
            st.dataframe(resumen_dist, hide_index=True, use_container_width=True)

        st.markdown("</div>", unsafe_allow_html=True)

        st.markdown("<div class='section-card'>", unsafe_allow_html=True)
        st.write("**Leyenda**")
        render_status_legend()
        st.markdown("</div>", unsafe_allow_html=True)

    with ms1:
        m_supervision = make_supervision_cluster_map(filtered, selected_tile=supervision_tile)
        st_folium(m_supervision, width=1280, height=720, returned_objects=[])


with tab_summary:
    st.subheader("Resumen operativo")

    if df.empty:
        st.warning("La base está vacía. Carga un Excel desde la pestaña Cargar Excel.")
    else:
        a, b = st.columns([1, 1])

        with a:
            st.write("**Registros por distrito y sección**")

            dist = (
                filtered.groupby(["distrito_local", "seccion"], dropna=False)
                .size()
                .reset_index(name="total")
            )

            st.dataframe(
                dist.sort_values(["distrito_local", "seccion"]),
                hide_index=True,
                use_container_width=True,
            )

        with b:
            st.write("**Estatus de supervisión**")

            status_df = filtered["estatus"].value_counts().rename_axis("estatus").reset_index(name="total")
            st.dataframe(status_df, hide_index=True, use_container_width=True)

            st.write("**Avance general**")
            avance = (total_verificados / total_registros * 100) if total_registros else 0
            st.progress(avance / 100)
            st.caption(f"{avance:.1f}% verificado sobre registros totales.")

        st.write("**Estado de coordenadas**")

        coord_df = (
            df["estado_coordenada"]
            .fillna("pendiente")
            .replace("", "pendiente")
            .value_counts()
            .rename_axis("estado_coordenada")
            .reset_index(name="total")
        )

        st.dataframe(coord_df, hide_index=True, use_container_width=True)


with tab_review:
    st.subheader("Revisión individual")

    mapped_filtered = filtered.dropna(subset=["latitud_mapa", "longitud_mapa"]) if not filtered.empty else filtered

    if mapped_filtered.empty:
        st.warning("No hay registros mapeables con los filtros actuales.")
    else:
        filtered_options = mapped_filtered.sort_values(["distrito_local", "seccion", "id"]).copy()

        option_labels = {
            int(row.id): f"ID {int(row.id)} | D{row.distrito_local} S{row.seccion} | {row.colonia} | {str(row.direccion)[:55]}"
            for _, row in filtered_options.iterrows()
        }

        default_id = st.session_state.get("selected_lona_id")
        ids = list(option_labels.keys())
        default_index = ids.index(default_id) if default_id in ids else 0

        selected_id = st.selectbox(
            "Selecciona el registro a revisar",
            ids,
            index=default_index,
            format_func=lambda x: option_labels.get(x, str(x)),
        )

        st.session_state["selected_lona_id"] = selected_id

        row = df[df["id"].astype(int) == int(selected_id)].iloc[0]

        left, right = st.columns([1.15, 1])

        with left:
            st.markdown(f"### ID {selected_id}")
            st.write(f"**Distrito local:** {row.get('distrito_local','')}  |  **Sección:** {row.get('seccion','')}")
            st.write(f"**Colonia:** {row.get('colonia','')}")
            st.write(f"**Dirección:** {row.get('direccion','')}")
            st.write(f"**Responsable:** {row.get('responsable','')}")
            st.write(f"**Observaciones:** {row.get('observaciones','')}")
            st.write(f"**Fuente coordenada:** {row.get('fuente_coordenada','')}")

            if clean_text(row.get("link_maps", "")):
                st.link_button("Abrir link original de Google Maps", clean_text(row.get("link_maps")))

            imgs = image_files_for_lona(row)

            if imgs:
                st.markdown("**Evidencia fotográfica**")

                for img in imgs[:3]:
                    st.image(str(img), caption=img.name, use_container_width=True)
            else:
                st.info("No se encontró imagen de evidencia para este registro.")

        with right:
            with st.form("form_revision"):
                estatus_actual = clean_text(row.get("estatus", "Pendiente")) or "Pendiente"

                estatus = st.selectbox(
                    "Estatus de supervisión",
                    STATUS_OPTIONS,
                    index=STATUS_OPTIONS.index(estatus_actual) if estatus_actual in STATUS_OPTIONS else 0,
                )

                supervisor = st.text_input(
                    "Supervisor",
                    value=clean_text(row.get("supervisor", "")) or st.session_state.get("auth_name", ""),
                )

                nota = st.text_area(
                    "Nota de supervisión",
                    value=clean_text(row.get("nota_supervision", "")),
                    height=120,
                )

                st.markdown("**Corrección opcional de coordenadas**")
                st.caption("Déjalas vacías si la ubicación original es correcta.")

                lat_corr = st.text_input("Latitud corregida")
                lon_corr = st.text_input("Longitud corregida")

                guardar = st.form_submit_button("Guardar revisión", use_container_width=True)

                if guardar:
                    lat_val = to_float(lat_corr)
                    lon_val = to_float(lon_corr)

                    if (clean_text(lat_corr) and lat_val is None) or (clean_text(lon_corr) and lon_val is None):
                        st.error("Las coordenadas corregidas no son válidas.")
                    elif (lat_val is not None and not -90 <= lat_val <= 90) or (lon_val is not None and not -180 <= lon_val <= 180):
                        st.error("Las coordenadas corregidas están fuera de rango.")
                    else:
                        update_lona_review(
                            selected_id,
                            estatus,
                            supervisor.strip(),
                            nota.strip(),
                            lat_val,
                            lon_val,
                        )
                        st.success("Revisión guardada.")
                        st.rerun()

            st.markdown("**Vista rápida del punto**")

            mini = load_lonas_df()
            mini = mini[mini["id"].astype(int) == int(selected_id)]

            st_folium(
                make_map(
                    mini,
                    selected_id,
                    selected_tile=map_style,
                    cluster_points=False,
                ),
                width=650,
                height=320,
                returned_objects=[],
            )


with tab_table:
    st.subheader("Tabla filtrada")

    if filtered.empty:
        st.warning("No hay registros con los filtros actuales.")
    else:
        show_cols = [
            "id",
            "archivo_origen",
            "fecha",
            "responsable",
            "municipio",
            "distrito_local",
            "seccion",
            "colonia",
            "direccion",
            "nombre_enlace",
            "celular",
            "lonas_colocadas",
            "estatus",
            "supervisor",
            "nota_supervision",
            "fecha_revision",
            "latitud_mapa",
            "longitud_mapa",
            "fuente_coordenada",
            "estado_coordenada",
            "observaciones",
            "link_maps",
        ]

        show_cols = [c for c in show_cols if c in filtered.columns]

        st.dataframe(
            filtered[show_cols].sort_values(["distrito_local", "seccion", "id"]),
            use_container_width=True,
            hide_index=True,
        )


with tab_pending:
    st.subheader("Pendientes sin coordenada")

    pendientes_df = df[df["latitud_mapa"].isna() | df["longitud_mapa"].isna()].copy() if not df.empty else pd.DataFrame()

    st.caption(
        "Estos registros no tienen coordenada exacta. Puedes intentar resolver links cortos o capturar coordenadas manualmente."
    )

    if pendientes_df.empty:
        st.success("No hay pendientes sin coordenada.")
    else:
        c1, c2, c3 = st.columns([1, 1, 1])

        with c1:
            block_size = st.number_input(
                "Resolver por bloque",
                min_value=1,
                max_value=500,
                value=50,
                step=10,
            )

        with c2:
            st.write("")
            st.write("")
            resolve_block = st.button("Resolver links cortos pendientes", use_container_width=True)

        with c3:
            st.download_button(
                "Descargar pendientes CSV",
                dataframe_to_csv_bytes(pendientes_df),
                file_name="lonas_pendientes_sin_coordenada.csv",
                mime="text/csv",
                use_container_width=True,
            )

        if resolve_block:
            to_resolve = pendientes_df[pendientes_df["link_maps"].astype(str).str.len() > 0].head(int(block_size)).copy()

            if to_resolve.empty:
                st.warning("No hay links para resolver en el bloque seleccionado.")
            else:
                progress = st.progress(0)
                ok_count = 0

                for i, (_, row) in enumerate(to_resolve.iterrows(), start=1):
                    result = resolve_maps_link(row.get("link_maps", ""), try_expand=True)

                    if result["latitud"] is not None and result["longitud"] is not None:
                        update_lona_coords(
                            int(row["id"]),
                            result["latitud"],
                            result["longitud"],
                            result["fuente_coordenada"],
                            result["estado_coordenada"],
                            result["url_expandida"],
                        )
                        ok_count += 1
                    else:
                        update_lona_expanded_url(int(row["id"]), result.get("url_expandida", ""))

                    progress.progress(i / len(to_resolve))

                st.success(f"Proceso terminado. Coordenadas resueltas: {ok_count} de {len(to_resolve)}.")
                st.rerun()

        st.write("### Captura manual")

        pending_options = {
            int(row.id): f"ID {int(row.id)} | D{row.distrito_local} S{row.seccion} | {row.colonia} | {str(row.direccion)[:55]}"
            for _, row in pendientes_df.head(1000).iterrows()
        }

        selected_pending = st.selectbox(
            "Selecciona pendiente",
            list(pending_options.keys()),
            format_func=lambda x: pending_options.get(x, str(x)),
        )

        row_p = pendientes_df[pendientes_df["id"].astype(int) == int(selected_pending)].iloc[0]

        st.write(f"**Dirección:** {row_p.get('direccion','')}")
        st.write(f"**Colonia:** {row_p.get('colonia','')}")
        st.write(f"**Responsable:** {row_p.get('responsable','')}")

        if clean_text(row_p.get("link_maps", "")):
            st.link_button("Abrir link original", clean_text(row_p.get("link_maps")))

        with st.form("manual_coords_form"):
            lat_manual = st.text_input("Latitud")
            lon_manual = st.text_input("Longitud")
            save_manual = st.form_submit_button("Guardar coordenada manual", use_container_width=True)

            if save_manual:
                lat = to_float(lat_manual)
                lon = to_float(lon_manual)

                if lat is None or lon is None:
                    st.error("Captura latitud y longitud válidas.")
                elif not (-90 <= lat <= 90 and -180 <= lon <= 180):
                    st.error("Las coordenadas están fuera de rango.")
                else:
                    update_lona_coords(
                        int(selected_pending),
                        lat,
                        lon,
                        "captura_manual",
                        "exacta",
                        "",
                    )
                    st.success("Coordenada guardada.")
                    st.rerun()

        st.write("### Tabla de pendientes")

        pending_cols = [
            "id",
            "archivo_origen",
            "responsable",
            "municipio",
            "distrito_local",
            "seccion",
            "colonia",
            "direccion",
            "nombre_enlace",
            "celular",
            "link_maps",
            "url_expandida",
            "fuente_coordenada",
            "estado_coordenada",
            "observaciones",
        ]

        pending_cols = [c for c in pending_cols if c in pendientes_df.columns]

        st.dataframe(
            pendientes_df[pending_cols],
            use_container_width=True,
            hide_index=True,
        )


with tab_export:
    st.subheader("Exportaciones")

    st.write("Descarga los avances de supervisión y genera un KMZ actualizado con los estatus y coordenadas corregidas.")

    export_all = st.checkbox("Exportar todos los registros", value=True)
    export_df = df if export_all else filtered

    c1, c2, c3, c4 = st.columns(4)

    with c1:
        st.download_button(
            "CSV de supervisión",
            dataframe_to_csv_bytes(export_df),
            file_name="supervision_lonas.csv",
            mime="text/csv",
            use_container_width=True,
        )

    with c2:
        kmz_bytes = build_kmz_bytes(export_df)

        st.download_button(
            "KMZ actualizado",
            kmz_bytes,
            file_name="lonas_supervision_actualizado.kmz",
            mime="application/vnd.google-earth.kmz",
            use_container_width=True,
        )

    with c3:
        if DB_FILE.exists():
            st.download_button(
                "Base SQLite",
                DB_FILE.read_bytes(),
                file_name="lonas_supervision.db",
                mime="application/octet-stream",
                use_container_width=True,
            )

    with c4:
        json_bytes = export_df.to_json(orient="records", force_ascii=False, indent=2).encode("utf-8")

        st.download_button(
            "JSON",
            json_bytes,
            file_name="lonas_supervision.json",
            mime="application/json",
            use_container_width=True,
        )

    st.info(
        "En Streamlit Cloud la base SQLite local puede reiniciarse. "
        "Para operación multiusuario formal conviene migrar después a Supabase."
    )


with tab_help:
    st.subheader("Guía rápida de uso")

    st.markdown(
        """
        **Flujo recomendado:**

        1. Entra con usuario y contraseña.
        2. Carga nuevos Excel desde la pestaña **Cargar Excel**.
        3. Evita resolver 1,000+ links al cargar; primero guarda la base.
        4. En **Pendientes sin coordenada**, resuelve links por bloques.
        5. Captura coordenadas manuales cuando el link corto no se pueda resolver.
        6. Revisa cada punto desde **Supervisión**.
        7. Exporta CSV, KMZ o base SQLite desde **Exportar**.

        **Estatus sugeridos:**

        - **Pendiente:** aún no revisado.
        - **Verificado:** ubicación y evidencia correctas.
        - **Corregir ubicación:** se detectó que el punto requiere ajuste.
        - **Retirar/Reponer lona:** hay incidencia física con la lona.
        - **No localizada:** no se encontró en campo.

        **Usuarios:** se agregan en Streamlit Secrets, no dentro del código.
        """
    )


with tab_upload:
    st.subheader("Cargar Excel")

    st.caption(
        "Carga nuevos reportes con formato de lonas. "
        "La app detecta encabezados, guarda en SQLite y separa mapeables/pendientes."
    )

    uploaded_files = st.file_uploader(
        "Sube uno o varios archivos Excel",
        type=["xlsx", "xlsm", "xls"],
        accept_multiple_files=True,
    )

    resolve_on_upload = st.checkbox(
        "Intentar resolver links cortos al cargar",
        value=False,
        help="Para archivos grandes conviene dejarlo desactivado y resolver por bloques desde Pendientes sin coordenada.",
    )

    if uploaded_files:
        st.write("**Archivos recibidos:**")

        for f in uploaded_files:
            st.write(f"- {f.name}")

        if st.button("Procesar e integrar a la base SQLite", use_container_width=True):
            results = []

            with st.spinner("Procesando archivos..."):
                for uploaded_file in uploaded_files:
                    try:
                        total, inserted, duplicated, mapeables, pendientes = parse_excel_file(
                            uploaded_file,
                            resolve_short_links=resolve_on_upload,
                        )

                        results.append(
                            {
                                "archivo": uploaded_file.name,
                                "registros_leidos": total,
                                "insertados": inserted,
                                "duplicados": duplicated,
                                "mapeables": mapeables,
                                "pendientes": pendientes,
                            }
                        )

                    except Exception as e:
                        results.append(
                            {
                                "archivo": uploaded_file.name,
                                "error": str(e),
                            }
                        )

            st.success("Proceso terminado.")
            st.dataframe(pd.DataFrame(results), hide_index=True, use_container_width=True)
            st.rerun()

    st.divider()
    st.subheader("Vista de base actual")

    st.write(f"Registros totales en SQLite: **{total_registros}**")
    st.write(f"Mapeables: **{total_mapeables}**")
    st.write(f"Pendientes sin coordenada: **{total_pendientes_coord}**")

    with get_conn() as conn:
        uploads_df = pd.read_sql_query(
            "SELECT * FROM uploads ORDER BY id DESC LIMIT 20",
            conn,
        )

    if not uploads_df.empty:
        st.write("**Últimas cargas:**")
        st.dataframe(uploads_df, hide_index=True, use_container_width=True)
